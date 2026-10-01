"""
Local Document-to-Data Parsing Engine for Reva (100% Free / Zero SaaS Fees).
Extracts text from PDF, DOCX, PPTX, and plain text, then normalizes into Trium's IdeaBrief schema.
"""

import os
from typing import Any, Dict, Optional
from pydantic import BaseModel, Field

import pdfplumber
from docx import Document
from pptx import Presentation

from src.utils.llm_client import llm_client


class IdeaBrief(BaseModel):
    name: str = Field(description="Name or title of the venture idea")
    sector: Optional[str] = Field(default=None, description="Industry sector or domain")
    description: str = Field(description="Brief 2-3 sentence overview of the business model")
    problem: str = Field(description="Specific customer/partner pain point being addressed")
    solution: str = Field(description="Proposed product, service, or technology solution")
    targetCustomer: str = Field(description="Target audience or user segment")
    goToMarket: Optional[str] = Field(default=None, description="Distribution channels and customer acquisition strategy")
    monetization: Optional[str] = Field(default=None, description="Revenue model, pricing, or unit economics")
    similarSolutions: Optional[str] = Field(default=None, description="Known competitors or informal alternatives")


EXTRACTION_SYSTEM_PROMPT = """You are a senior venture architect at Trium Limited, an African venture studio.
Extract structured venture concept details from the provided document text.
Return ONLY a valid JSON object matching this schema:
{
  "name": "<short distinctive name>",
  "sector": "<industry sector, e.g. Fintech, Agri-Logistics, CleanTech>",
  "description": "<concise 2-3 sentence overview>",
  "problem": "<specific acute problem being solved>",
  "solution": "<proposed technical & operational solution>",
  "targetCustomer": "<target customer profile and market segment>",
  "goToMarket": "<go-to-market channels or null>",
  "monetization": "<revenue model, pricing, or null>",
  "similarSolutions": "<competitors or informal alternatives or null>"
}
Do not hallucinate facts. If a field is not present in the document, set it to null."""


class DocumentParser:
    """Extracts raw text from local files and parses into IdeaBrief using free LLM."""

    @staticmethod
    def extract_text_from_file(file_path: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return DocumentParser._extract_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return DocumentParser._extract_docx(file_path)
        elif ext in [".pptx", ".ppt"]:
            return DocumentParser._extract_pptx(file_path)
        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def _extract_pdf(file_path: str) -> str:
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page_idx, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- Page {page_idx + 1} ---\n{page_text}")
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_docx(file_path: str) -> str:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)

    @staticmethod
    def _extract_pptx(file_path: str) -> str:
        prs = Presentation(file_path)
        slide_texts = []
        for idx, slide in enumerate(prs.slides):
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_content.append(text)
            if slide_content:
                slide_texts.append(f"--- Slide {idx + 1} ---\n" + "\n".join(slide_content))
        return "\n\n".join(slide_texts)

    def parse_to_brief(self, raw_input: str, is_file_path: bool = False) -> IdeaBrief:
        """
        Takes either raw text prompt or file path and structures it into IdeaBrief.
        """
        if is_file_path:
            text = self.extract_text_from_file(raw_input)
        else:
            text = raw_input.strip()

        if not text:
            raise ValueError("Input content is empty.")

        # Limit context to first 25,000 chars for prompt efficiency
        prompt = f"Document content:\n{text[:25000]}\n\nExtract the structured venture idea details."
        data = llm_client.generate_json(prompt, system_instruction=EXTRACTION_SYSTEM_PROMPT)

        return IdeaBrief(
            name=data.get("name") or "Unnamed Initiative",
            sector=data.get("sector") or "General Tech",
            description=data.get("description") or (text[:300] + "..."),
            problem=data.get("problem") or "Problem statement not explicitly specified.",
            solution=data.get("solution") or "Solution mechanics not explicitly specified.",
            targetCustomer=data.get("targetCustomer") or "Target segment not explicitly specified.",
            goToMarket=data.get("goToMarket"),
            monetization=data.get("monetization"),
            similarSolutions=data.get("similarSolutions"),
        )


doc_parser = DocumentParser()
